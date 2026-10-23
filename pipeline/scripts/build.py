"""
Build everything people look at from the master database (the store folder + facility_extras.xlsx):
  out/RI_Nursing_Homes_Database.xlsx   the Excel workbook
  out/dashboard/data.js                the data file the dashboard page reads

Usage:  python build.py            (rebuild only, e.g. after editing facility_extras.xlsx)
"""
import json
import os
import re
import sys
from datetime import datetime

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

# ------------------------------------------------------------------ metric dictionary
# key, store column, label, better ("up" = higher is better for residents, "down" = lower is better, "" = neither)
METRICS = [
    ("overall", "overall_rating", "Overall star rating", "up"),
    ("health", "health_inspection_rating", "Health inspection rating", "up"),
    ("qm", "qm_rating", "Quality measure rating", "up"),
    ("staffing", "staffing_rating", "Staffing rating", "up"),
    ("hprd", "reported_total_nurse_staffing_hours_per_resident_per_day", "Total nurse staff hours per resident per day", "up"),
    ("rn_hprd", "reported_rn_staffing_hours_per_resident_per_day", "RN hours per resident per day", "up"),
    ("lpn_hprd", "reported_lpn_staffing_hours_per_resident_per_day", "LPN hours per resident per day", "up"),
    ("aide_hprd", "reported_nurse_aide_staffing_hours_per_resident_per_day", "Nurse aide hours per resident per day", "up"),
    ("wkend_hprd", "total_number_of_nurse_staff_hours_per_resident_per_day_on_the_weekend", "Weekend nurse staff hours per resident per day", "up"),
    ("cm_hprd", "case_mix_total_nurse_staffing_hours_per_resident_per_day", "Expected nurse hours given resident needs (case-mix)", ""),
    ("adj_hprd", "adjusted_total_nurse_staffing_hours_per_resident_per_day", "Adjusted nurse staff hours per resident per day", "up"),
    ("cmi", "nursing_case_mix_index", "Nursing case-mix index (higher = residents need more care)", ""),
    ("turnover", "total_nursing_staff_turnover", "Nursing staff turnover (%)", "down"),
    ("rn_turnover", "registered_nurse_turnover", "RN turnover (%)", "down"),
    ("admin_left", "number_of_administrators_who_have_left_the_nursing_home", "Administrators who left (past year)", "down"),
    ("residents", "average_number_of_residents_per_day", "Average residents per day", ""),
    ("beds", "number_of_certified_beds", "Certified beds", ""),
    ("health_score", "total_weighted_health_survey_score", "Weighted health inspection score (lower is better)", "down"),
    ("cycle1_defs", "rating_cycle_1_total_number_of_health_deficiencies", "Health citations, most recent inspection cycle", "down"),
    ("fines_n", "number_of_fines", "Number of fines (CMS 3-year window)", "down"),
    ("fines_amt", "total_amount_of_fines_in_dollars", "Total fines, $ (CMS 3-year window)", "down"),
    ("subst_complaints", "number_of_substantiated_complaints", "Substantiated complaints (CMS 3-year window)", "down"),
    ("fri", "number_of_facility_reported_incidents", "Facility-reported incidents (CMS 3-year window)", "down"),
]
NUMERIC = [m[1] for m in METRICS]
TREND_SHEETS = [("overall", "Trend Overall"), ("health", "Trend Inspections"), ("staffing", "Trend Staffing Rating"),
                ("qm", "Trend Quality Rating"), ("hprd", "Trend Nurse Hours"), ("rn_hprd", "Trend RN Hours"),
                ("turnover", "Trend Turnover")]

SEVERITY = {**{k: "Minimal harm potential" for k in "ABC"},
            **{k: "Potential for more than minimal harm" for k in "DEF"},
            **{k: "Actual harm" for k in "GHI"},
            **{k: "Immediate jeopardy" for k in "JKL"}}

SCOPE = {**{k: "Isolated (one or a few residents)" for k in "ADGJ"},
         **{k: "Pattern (more than a few residents)" for k in "BEHK"},
         **{k: "Widespread (many residents)" for k in "CFIL"}}
CC_PDF = "https://www.medicare.gov/care-compare/inspections/pdf/nursing-home/{ccn}/health/{kind}?date={date}"
CC_PAGE = "https://www.medicare.gov/care-compare/details/nursing-home/{ccn}/view-all/?state=RI"
PROPUBLICA = "https://projects.propublica.org/nursing-homes/homes/h-{ccn}"

KEY_QMS = {  # MDS and claims measures shown on the dashboard
    "481": "Long-stay residents given antipsychotic medication (%)",
    "410": "Long-stay residents with falls with major injury (%)",
    "479": "Long-stay residents with pressure ulcers (%)",
    "407": "Long-stay residents with a urinary tract infection (%)",
    "404": "Long-stay residents who lose too much weight (%)",
    "401": "Long-stay residents whose need for help with daily activities increased (%)",
    "452": "Long-stay residents given antianxiety or sleep medication (%)",
    "408": "Long-stay residents with depressive symptoms (%)",
    "434": "Short-stay residents newly given antipsychotic medication (%)",
    "551": "Hospitalizations per 1,000 long-stay resident days",
    "552": "ER visits per 1,000 long-stay resident days",
    "521": "Short-stay residents rehospitalized (%)",
}



MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"]


def closed_from_note(note):
    """A facility details note that starts "Closed <Month> <Year>" marks a home closed even if CMS still lists it.
    Returns the month as YYYY-MM, or None."""
    m = re.match(r"\s*closed\s+([a-z]+)\s+(\d{4})", str(note or ""), re.I)
    if m and m.group(1).lower() in MONTHS:
        return f"{m.group(2)}-{MONTHS.index(m.group(1).lower()) + 1:02d}"
    return None

def month_label(m):
    return datetime.strptime(m, "%Y-%m").strftime("%b %Y")


def norm_name(s):
    s = str(s).upper().replace("&", " AND ")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\b(LLC|L L C|INC|THE|CORP|CORPORATION|LTD|CO|DBA)\b", " ", s)
    s = re.sub(r"\bCENTRE\b", "CENTER", s)
    return " ".join(s.split())


def num(s):
    return pd.to_numeric(s, errors="coerce")


# ------------------------------------------------------------------ load
def load(root):
    st = os.path.join(root, "store")
    rd = lambda n: pd.read_csv(os.path.join(st, n + ".csv"), dtype=str, keep_default_na=False, na_values=[""])
    d = {n: rd(n) for n in ["provider_monthly", "health_citations", "fire_citations", "penalties", "ownership",
                            "qm_mds", "qm_claims", "measures", "state_us_averages", "survey_dates", "tags",
                            "releases_loaded"]}
    hist_p = os.path.join(st, "kg_2017_2018_ratings.csv")
    d["hist"] = pd.read_csv(hist_p, dtype=str) if os.path.exists(hist_p) else None
    d["extras"] = pd.read_excel(os.path.join(root, "facility_extras.xlsx"), sheet_name="Extras", dtype=str).fillna("")
    d["extras"]["ccn"] = d["extras"]["ccn"].str.replace(r"\.0$", "", regex=True)
    return d


# ------------------------------------------------------------------ derived tables
def derive(d):
    pm = d["provider_monthly"].copy()
    for c in NUMERIC:
        pm[c] = num(pm[c])
    if d["hist"] is not None:
        h = d["hist"].copy()
        for c in h.columns:
            if c not in ("ccn", "cms_month"):
                h[c] = num(h[c])
        h["source_note"] = "From kg's earlier workbook"
        pm["source_note"] = ""
        pm = pd.concat([h, pm], ignore_index=True)
    else:
        pm["source_note"] = ""
    pm = pm.sort_values(["ccn", "cms_month"]).reset_index(drop=True)
    pm["occupancy_pct"] = (pm["average_number_of_residents_per_day"] / pm["number_of_certified_beds"] * 100).round(1)

    months = sorted(d["provider_monthly"]["cms_month"].unique())
    latest = months[-1]
    all_months = sorted(pm["cms_month"].unique())

    ex = d["extras"]
    name_of = dict(zip(ex["ccn"], ex["display_name"]))
    cms_rows = pm[pm["source_note"].fillna("") == ""] if "source_note" in pm.columns else pm
    last = cms_rows.groupby("ccn").tail(1).set_index("ccn")
    first = cms_rows.groupby("ccn").head(1).set_index("ccn")

    # ---------- name and operator history
    hist_rows = []
    for ccn, g in cms_rows.groupby("ccn"):
        for col, kind in (("provider_name", "Facility name"), ("legal_business_name", "Legal business (operator) name")):
            vals = g[["cms_month", col]].copy()
            vals[col] = vals[col].where(~vals[col].fillna("").str.contains("NOT AVAILABLE", case=False))
            vals[col] = vals[col].ffill()
            vals["k"] = vals[col].map(norm_name)
            vals["grp"] = (vals["k"] != vals["k"].shift()).cumsum()
            for _, gg in vals.dropna(subset=[col]).groupby("grp"):
                hist_rows.append({"ccn": ccn, "kind": kind, "name": gg[col].iloc[-1],
                                  "from": gg["cms_month"].min(), "to": gg["cms_month"].max()})
    names = pd.DataFrame(hist_rows)
    chow = cms_rows[cms_rows["provider_changed_ownership_in_last_12_months"] == "Y"].groupby("ccn")["cms_month"].agg(["min", "max"])

    # ---------- citations
    hc = d["health_citations"].copy(); hc["inspection"] = "Health"
    fc = d["fire_citations"].copy(); fc["inspection"] = "Fire safety"
    cit = pd.concat([hc, fc], ignore_index=True)
    cit["severity"] = cit["scope_severity_code"].map(SEVERITY)
    cit["scope"] = cit["scope_severity_code"].map(SCOPE)
    # Official CMS-2567 inspection report (Care Compare PDF). Care Compare only keeps reports for the
    # inspections it currently lists, so link only citations in the latest release.
    def _report(r):
        if r["inspection"] != "Health" or r["last_seen"] != latest or not isinstance(r["survey_date"], str):
            return ""
        kind = "health-inspection" if r["standard_deficiency"] == "Y" else (
            "complaint-inspection" if r["complaint_deficiency"] == "Y" else "")
        return CC_PDF.format(ccn=r["ccn"], kind=kind, date=r["survey_date"]) if kind else ""
    cit["report_url"] = cit.apply(_report, axis=1)
    cit["year"] = cit["survey_date"].str[:4]
    cit["harm"] = cit["scope_severity_code"].isin(list("GHIJKL"))
    cit["ij"] = cit["scope_severity_code"].isin(list("JKL"))
    cit["complaint"] = cit["complaint_deficiency"] == "Y"
    cit["infection_control"] = cit["infection_control_inspection_deficiency"] == "Y"
    cit["tag"] = cit["deficiency_prefix"].fillna("") + cit["deficiency_tag_number"].fillna("")
    cit["listed_now"] = cit["last_seen"] == latest
    cit["cms_status"] = np.where(cit["listed_now"], "Listed in latest CMS release", "No longer listed by CMS")
    cit = cit.sort_values(["survey_date", "ccn"], ascending=[False, True])

    hcy = cit[cit["inspection"] == "Health"].groupby(["ccn", "year"]).agg(
        health_citations=("tag", "size"), complaint_citations=("complaint", "sum"),
        harm_citations=("harm", "sum"), immediate_jeopardy=("ij", "sum"),
        infection_control=("infection_control", "sum")).reset_index()
    fcy = cit[cit["inspection"] == "Fire safety"].groupby(["ccn", "year"]).size().rename("fire_safety_citations").reset_index()
    cy = hcy.merge(fcy, on=["ccn", "year"], how="outer").fillna(0)
    for c in cy.columns[2:]:
        cy[c] = cy[c].astype(int)

    # ---------- penalties
    pen = d["penalties"].copy()
    pen["fine_amount"] = num(pen["fine_amount"])
    pen["payment_denial_length_in_days"] = num(pen["payment_denial_length_in_days"])
    pen["year"] = pen["penalty_date"].str[:4]
    # A fine that disappears from CMS data well before its 3-year window ends was revised (new date or amount)
    # or withdrawn; the replacement shows up as its own record, so leave these out of totals.
    rels = sorted(d["releases_loaded"]["cms_month"])
    nxt = {m: (rels[i + 1] if i + 1 < len(rels) else None) for i, m in enumerate(rels)}
    gone_in = pen["last_seen"].map(nxt)
    age = (pd.to_datetime(gone_in + "-01", errors="coerce") - pd.to_datetime(pen["penalty_date"], errors="coerce")).dt.days / 30.44
    pen["cms_status"] = np.where(pen["last_seen"] == latest, "Listed in latest CMS release",
                                 np.where(age < 34, "Removed by CMS early (likely revised or withdrawn); not counted",
                                          "Aged out of CMS's 3-year window"))
    pen["counted"] = ~pen["cms_status"].str.startswith("Removed")
    pen = pen.sort_values(["penalty_date"], ascending=False)
    py = pen[pen["counted"]].groupby(["ccn", "year"]).agg(fines=("penalty_type", lambda s: (s == "Fine").sum()),
                                          fine_dollars=("fine_amount", "sum"),
                                          payment_denials=("penalty_type", lambda s: (s == "Payment Denial").sum())).reset_index()

    # ---------- ownership
    own = d["ownership"].copy()
    own["current"] = own["last_seen"] == latest
    own = own.sort_values(["ccn", "current", "role", "owner_name"], ascending=[True, False, True, True])

    # ---------- 3-year window counts for the Current tab
    cutoff = (pd.Timestamp(latest + "-01") - pd.DateOffset(years=3)).strftime("%Y-%m-%d")
    rc = cit[(cit["survey_date"] >= cutoff) & (cit["inspection"] == "Health") & cit["listed_now"]]
    c3 = rc.groupby("ccn").agg(cit3=("tag", "size"), harm3=("harm", "sum"), ij3=("ij", "sum"),
                               compl3=("complaint", "sum"))
    p3 = pen[(pen["penalty_date"] >= cutoff) & (pen["last_seen"] == latest)].groupby("ccn").agg(
        fines3_n=("penalty_type", lambda s: (s == "Fine").sum()), fines3_amt=("fine_amount", "sum"),
        denials3=("penalty_type", lambda s: (s == "Payment Denial").sum()))
    last_std = d["survey_dates"]
    last_std = last_std[last_std["type_of_survey"].fillna("").str.contains("Health Standard", case=False)]
    last_std = last_std.groupby("ccn")["survey_date"].max()

    # ---------- facilities master
    fac = []
    ccns = list(dict.fromkeys(list(ex["ccn"][ex["ccn"] != ""]) + list(last.index)))
    for ccn in ccns:
        e = ex[ex["ccn"] == ccn].iloc[0].to_dict() if (ex["ccn"] == ccn).any() else {}
        if ccn in last.index:
            L = last.loc[ccn]
            status = "Open" if L["cms_month"] == latest else f"No longer in CMS data (last listed {month_label(L['cms_month'])})"
            closed_month = closed_from_note(e.get("notes"))   # "Closed January 2025 ..." in the facility details notes
            closed_last = None
            if closed_month:
                status = f"Closed {month_label(closed_month)}"
                # show the home as of the last CMS file on or before it closed (CMS skips some months)
                before = sorted(m for m in pm.loc[pm["ccn"] == ccn, "cms_month"].unique() if m <= closed_month)
                closed_last = before[-1] if before else L["cms_month"]
            rec = {"ccn": ccn, "name": e.get("display_name") or L["provider_name"], "cms_name": L["provider_name"],
                   "legal_business_name": L["legal_business_name"], "address": L["provider_address"],
                   "city": str(L["city_town"]).title(), "zip": L["zip_code"], "county": str(L["county_parish"]).title(),
                   "phone": L["telephone_number"], "status": status,
                   "first_month": first.loc[ccn, "cms_month"], "last_month": closed_last or L["cms_month"],
                   "beds": L["number_of_certified_beds"], "residents": L["average_number_of_residents_per_day"],
                   "ownership_type": L["ownership_type"], "cms_chain": L["chain_name"],
                   "chain_facilities": L["number_of_facilities_in_chain"],
                   "ccrc": L["continuing_care_retirement_community"], "in_hospital": L["provider_resides_in_hospital"],
                   "council": L["with_a_resident_and_family_council"],
                   "certified_since": L["date_first_approved_to_provide_medicare_and_medicaid_services"],
                   "lat": num(pd.Series([L["latitude"]])).iloc[0], "lon": num(pd.Series([L["longitude"]])).iloc[0],
                   "changed_ownership_flag_months": (f"{month_label(chow.loc[ccn, 'min'])} to {month_label(chow.loc[ccn, 'max'])}"
                                                     if ccn in chow.index else "")}
        else:
            rec = {"ccn": ccn, "name": e.get("display_name", ""), "status": "Not in CMS data"}
        for k, v in e.items():
            if k not in ("ccn", "display_name"):
                rec["x_" + k] = v
        fac.append(rec)
    for _, e in ex[ex["ccn"] == ""].iterrows():
        rec = {"ccn": "", "name": e["display_name"], "status": "Not in CMS data"}
        for k, v in e.items():
            if k not in ("ccn", "display_name"):
                rec["x_" + k] = v
        fac.append(rec)
    fac = pd.DataFrame(fac)
    former = names[names["kind"] == "Facility name"].copy()
    former["k"] = former["name"].map(norm_name)
    fn = {}
    for ccn, g in former.groupby("ccn"):
        cur = norm_name(last.loc[ccn, "provider_name"])
        seen, lst = set(), []
        for _, r in g.sort_values("from").iterrows():
            if r["k"] != cur and r["k"] not in seen:
                seen.add(r["k"]); lst.append(r["name"])
        fn[ccn] = "; ".join(lst)
    fac["former_names"] = fac["ccn"].map(fn).fillna("")
    fac["sort"] = fac["status"].map(lambda s: 0 if s == "Open" else (1 if s.startswith(("No longer", "Closed")) else 2))
    fac = fac.sort_values(["sort", "name"]).drop(columns="sort").reset_index(drop=True)
    name_of = dict(zip(fac["ccn"], fac["name"]))

    # ---------- current snapshot
    cur = pm[pm["cms_month"] == latest].set_index("ccn")

    def ago(n):
        m = (pd.Timestamp(latest + "-01") - pd.DateOffset(months=n)).strftime("%Y-%m")
        # use the closest available month at or before m
        avail = [x for x in all_months if x <= m]
        return pm[pm["cms_month"] == avail[-1]].set_index("ccn") if avail else None, (avail[-1] if avail else None)

    y1, m1 = ago(12)
    y3, m3 = ago(36)
    avg = d["state_us_averages"].copy()
    for c in avg.columns[2:]:
        avg[c] = num(avg[c])
    ri_now = avg[(avg["area"] == "RI") & (avg["cms_month"] == latest)]
    current = []
    for ccn in fac[fac["status"] == "Open"]["ccn"]:
        r = cur.loc[ccn]
        f = fac[fac["ccn"] == ccn].iloc[0]
        rec = {
            "CCN": ccn, "Facility": f["name"], "City": f["city"], "Owner group (ABC-RI)": f.get("x_owner_group", ""),
            "CMS chain": f["cms_chain"], "Ownership type": f["ownership_type"], "Beds": r["number_of_certified_beds"],
            "Residents per day": r["average_number_of_residents_per_day"], "Occupancy %": r["occupancy_pct"],
            "Overall rating": r["overall_rating"],
            f"Overall rating {month_label(m1)}": y1["overall_rating"].get(ccn) if y1 is not None else None,
            f"Overall rating {month_label(m3)}": y3["overall_rating"].get(ccn) if y3 is not None else None,
            "Health inspection rating": r["health_inspection_rating"], "Staffing rating": r["staffing_rating"],
            "Quality measure rating": r["qm_rating"],
            "Nurse hours per resident per day": r["reported_total_nurse_staffing_hours_per_resident_per_day"],
            "RN hours per resident per day": r["reported_rn_staffing_hours_per_resident_per_day"],
            "Weekend nurse hours per resident per day": r["total_number_of_nurse_staff_hours_per_resident_per_day_on_the_weekend"],
            "Nursing case-mix index": r["nursing_case_mix_index"],
            "Nursing staff turnover %": r["total_nursing_staff_turnover"], "RN turnover %": r["registered_nurse_turnover"],
            "Administrators who left": r["number_of_administrators_who_have_left_the_nursing_home"],
            "Special Focus Facility": r["special_focus_status"] if isinstance(r["special_focus_status"], str) else "",
            "Abuse icon": "Yes" if r["abuse_icon"] == "Y" else "",
            "Last standard inspection": last_std.get(ccn, ""),
            "Health citations, 3 yrs": c3["cit3"].get(ccn, 0), "Complaint citations, 3 yrs": c3["compl3"].get(ccn, 0),
            "Harm citations (G-L), 3 yrs": c3["harm3"].get(ccn, 0), "Immediate jeopardy (J-L), 3 yrs": c3["ij3"].get(ccn, 0),
            "Fines, 3 yrs": p3["fines3_n"].get(ccn, 0), "Fines $, 3 yrs": p3["fines3_amt"].get(ccn, 0),
            "Payment denials, 3 yrs": p3["denials3"].get(ccn, 0),
        }
        for k in ["union", "pct_dementia_2021_ltcfocus", "pct_medicaid_2021_ltcfocus", "direct_care_per_diem_2022_eohhs", "year_sold"]:
            rec[{"union": "Union", "pct_dementia_2021_ltcfocus": "% dementia (2021)", "pct_medicaid_2021_ltcfocus": "% Medicaid (2021)",
                 "direct_care_per_diem_2022_eohhs": "Direct care per diem 2022", "year_sold": "Year sold (ABC-RI)"}[k]] = f.get("x_" + k, "")
        current.append(rec)
    current = pd.DataFrame(current)
    for c in ["% dementia (2021)", "% Medicaid (2021)", "Direct care per diem 2022"]:
        current[c] = num(current[c])

    # ---------- owner group summary (current homes)
    og = current.groupby(current["Owner group (ABC-RI)"].replace("", "(not set)")).agg(
        homes=("CCN", "size"), beds=("Beds", "sum"), avg_overall=("Overall rating", "mean"),
        avg_nurse_hours=("Nurse hours per resident per day", "mean"), avg_rn_hours=("RN hours per resident per day", "mean"),
        avg_turnover=("Nursing staff turnover %", "mean"), harm3=("Harm citations (G-L), 3 yrs", "sum"),
        ij3=("Immediate jeopardy (J-L), 3 yrs", "sum"), fines3=("Fines $, 3 yrs", "sum")).reset_index()
    og.columns = ["Owner group (ABC-RI)", "Homes", "Beds", "Avg overall rating", "Avg nurse hours/resident/day",
                  "Avg RN hours/resident/day", "Avg turnover %", "Harm citations, 3 yrs", "Immediate jeopardy, 3 yrs",
                  "Fines $, 3 yrs"]
    og["Harm citations per home"] = og["Harm citations, 3 yrs"] / og["Homes"]
    og = og.sort_values("Homes", ascending=False)

    # ---------- quality measures
    meas = d["measures"].drop_duplicates("measure_code").set_index("measure_code")
    qm = d["qm_mds"].copy(); qm["score"] = num(qm["score"])
    qc = d["qm_claims"].copy()
    for c in ["adjusted_score", "observed_score", "expected_score"]:
        qc[c] = num(qc[c])

    return dict(pm=pm, months=months, all_months=all_months, latest=latest, fac=fac, names=names, cit=cit, cy=cy,
                pen=pen, py=py, own=own, current=current, og=og, avg=avg, qm=qm, qc=qc, meas=meas, name_of=name_of,
                releases=d["releases_loaded"], tags=d["tags"], surveys=d["survey_dates"])


# ------------------------------------------------------------------ Excel
def write_excel(D, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    F = Font(name="Helvetica", size=10)
    FB = Font(name="Helvetica", size=10, bold=True)
    FT = Font(name="Helvetica", size=14, bold=True)
    YEL = PatternFill("solid", fgColor="FFF2A8")   # kg's convention: yellow = better for residents
    BLU = PatternFill("solid", fgColor="BDD7EE")   # blue = worse for residents
    wb = Workbook()
    name_of = D["name_of"]
    L = month_label(D["latest"])

    def sheet(title, df, widths=None, freeze="C2", fmt=None, wrap_header=True):
        ws = wb.create_sheet(title)
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in row])
        for c in ws[1]:
            c.font = FB
            c.alignment = Alignment(wrap_text=wrap_header, vertical="bottom")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = F
        for i, col in enumerate(df.columns, 1):
            w = (widths or {}).get(col)
            if w is None:
                w = 34 if col in ("Facility", "facility", "Owner name", "Description") else max(9, min(18, len(str(col)) * 0.55 + 4))
            ws.column_dimensions[get_column_letter(i)].width = w
            if fmt and col in fmt:
                for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for c in cell:
                        c.number_format = fmt[col]
        ws.row_dimensions[1].height = 42
        if freeze:
            ws.freeze_panes = freeze
        if len(df):
            ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        return ws

    def col_letter(df, col):
        return get_column_letter(list(df.columns).index(col) + 1)

    def good_bad(ws, df, cols, better):
        """kg's color convention: yellow = better for residents, blue = worse."""
        n = len(df) + 1
        for col in cols:
            if col not in df.columns:
                continue
            L_ = col_letter(df, col)
            rng = f"{L_}2:{L_}{n}"
            if better == "stars":
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=YEL))
                ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.5", "2"], fill=BLU))
            else:
                lo, hi = ("FFF2A8", "BDD7EE") if better == "down" else ("BDD7EE", "FFF2A8")
                ws.conditional_formatting.add(rng, ColorScaleRule(start_type="percentile", start_value=10, start_color=lo,
                                                                  mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                                                  end_type="percentile", end_value=90, end_color=hi))

    # README
    ws = wb.active
    ws.title = "README"
    rows = [
        ("Rhode Island Nursing Homes Database", None),
        (f"Latest CMS data: {L} release.  Built {datetime.now().strftime('%B %d, %Y')}.", None),
        ("", None),
        ("What is in each tab", "header"),
        ("Current", f"One row per open home with the {L} ratings, staffing, turnover, plus 3-year counts of citations and fines and your extra data. Start here."),
        ("Facilities", "Every home, including closed ones and the few CMS doesn't cover: address, beds, status, former names, owner group, and your extra data."),
        ("Owner Groups", "Current homes summarized by owner group (from the owner_group column in facility_extras.xlsx)."),
        ("Name History", "Each facility name and legal business (operator) name, with the months CMS showed it. A new operator name usually means a change of ownership."),
        ("Ownership", "Every owner, officer, and managing entity CMS has listed since 2019, with the first and last release that listed them. 'Current' = listed in the latest release."),
        ("Monthly", "The full history: one row per home per CMS release, Jan 2019 to now. Filter or pivot this for anything."),
        ("Trend tabs", "Overall, Inspections, Staffing Rating, Quality Rating, Nurse Hours, RN Hours, Turnover: one row per home, one column per month, newest first. Yellow = better for residents, blue = worse."),
        ("Citations", "Every health and fire safety citation in CMS data, newest first, with severity and whether it came from a complaint."),
        ("Citations by Year", "Counts per home per year: all, complaint, harm (G-L), immediate jeopardy (J-L), infection control, fire safety."),
        ("Fines", "Every fine and payment denial, newest first."),
        ("Fines by Year", "Fines and dollars per home per year."),
        ("Quality Measures", "Quarterly resident-outcome measures (antipsychotics, falls, pressure ulcers, etc.) per home."),
        ("Hospital Measures", "Hospitalization and ER measures from Medicare claims, per home and measure period."),
        ("RI vs US", "Rhode Island and national averages for each CMS release."),
        ("Releases", "Which CMS monthly releases are in the database."),
        ("", None),
        ("Things to know", "header"),
        ("Gaps", "CMS did not publish (or the history files don't include) every month. Missing: " +
         ", ".join(month_label(m) for m in pd.period_range(D["months"][0], D["latest"], freq="M").strftime("%Y-%m")
                   if m not in D["months"]) + "."),
        ("CMS look-back windows", "Citation, fine, complaint, and incident counts on the facility summary cover roughly the last 3 years, so they rise and fall as old events drop off. The Citations and Fines tabs keep every record ever seen, so use those for counts by year."),
        ("Older citations", "The 2019 releases reached back to about 2015, so citations before ~2016 are incomplete."),
        ("Removed records", "The database keeps every citation and fine CMS ever published, even after CMS stops listing it. 3-year counts use only what CMS lists in the latest release, so they match Care Compare. Fines CMS dropped well before their 3-year window ended were almost always re-issued with a new date or amount; they are marked 'Removed by CMS early' and left out of all totals."),
        ("Staffing", "Staffing hours come from payroll data (PBJ) and lag the release date by about 6 months. Turnover starts in 2022 and the case-mix index in mid-2024 because CMS started publishing them then."),
        ("Case-mix index", "Higher numbers mean residents need MORE nursing care (higher acuity)."),
        ("Complaints/incidents", "CMS stopped publishing the substantiated complaint and facility-reported incident counts in early 2026. Complaint citations (Citations tab) are still available."),
        ("Severity letters", "A-C: potential for minimal harm. D-F: potential for more than minimal harm. G-I: actual harm. J-L: immediate jeopardy. Within each group the later letter means more widespread."),
        ("Your data", "Edit facility_extras.xlsx (owner group, union, dementia %, etc.), then rebuild. Don't edit this workbook directly; it gets regenerated each month."),
        ("", None),
        ("Source", "CMS Care Compare nursing home data, monthly archives: https://data.cms.gov/provider-data/archived-data/nursing-homes"),
    ]
    for a, b in rows:
        ws.append([a, None if b == "header" else b])
        r = ws.max_row
        ws.cell(r, 1).font = FB if b == "header" or r == 1 else F
        ws.cell(r, 2).font = F
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
    ws["A1"].font = FT
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120

    # Current
    cur = D["current"]
    ws = sheet("Current", cur, widths={"Facility": 38, "City": 15, "Owner group (ABC-RI)": 22, "CMS chain": 22, "Ownership type": 18},
               fmt={"Nurse hours per resident per day": "0.00", "RN hours per resident per day": "0.00",
                    "Weekend nurse hours per resident per day": "0.00", "Nursing case-mix index": "0.00",
                    "Fines $, 3 yrs": "$#,##0", "Direct care per diem 2022": "$#,##0.00", "Occupancy %": "0"})
    good_bad(ws, cur, [c for c in cur.columns if "rating" in c.lower()], "stars")
    good_bad(ws, cur, ["Nurse hours per resident per day", "RN hours per resident per day", "Weekend nurse hours per resident per day"], "up")
    good_bad(ws, cur, ["Nursing staff turnover %", "RN turnover %", "Harm citations (G-L), 3 yrs", "Fines $, 3 yrs", "Health citations, 3 yrs"], "down")

    # Facilities
    fac = D["fac"].copy()
    fac.columns = [c.replace("x_", "") for c in fac.columns]
    rename = {"ccn": "CCN", "name": "Facility", "cms_name": "Name in CMS data", "legal_business_name": "Legal business name",
              "former_names": "Former names", "status": "Status", "first_month": "First in data", "last_month": "Last in data",
              "changed_ownership_flag_months": "CMS 'changed ownership' flag shown"}
    fac = fac.rename(columns=rename)
    order = ["CCN", "Facility", "Status", "city", "address", "zip", "county", "phone", "Former names", "Name in CMS data",
             "Legal business name", "owner_group", "ownership_category", "cms_chain", "ownership_type", "year_sold",
             "CMS 'changed ownership' flag shown", "beds", "residents", "certified_since", "First in data", "Last in data"]
    order += [c for c in fac.columns if c not in order and c not in ("lat", "lon")] + ["lat", "lon"]
    sheet("Facilities", fac[order], widths={"Facility": 38, "Former names": 40, "Name in CMS data": 30, "Legal business name": 30,
                                            "address": 26, "Status": 22, "owner_notes": 40})

    # Owner groups
    sheet("Owner Groups", D["og"].round(2), widths={"Owner group (ABC-RI)": 34}, freeze="B2",
          fmt={"Fines $, 3 yrs": "$#,##0", "Harm citations per home": "0.0"})

    # Name history
    nh = D["names"].copy()
    nh.insert(1, "Facility", nh["ccn"].map(name_of))
    nh["from"] = nh["from"].map(month_label); nh["to"] = nh["to"].map(month_label)
    nh.columns = ["CCN", "Facility", "Type", "Name", "First shown", "Last shown"]
    sheet("Name History", nh, widths={"Facility": 34, "Type": 26, "Name": 50})

    # Ownership
    own = D["own"].copy()
    own.insert(1, "Facility", own["ccn"].map(name_of))
    own = own[["ccn", "Facility", "current", "role", "owner_name", "owner_type", "ownership_percentage", "since", "first_seen", "last_seen"]]
    own["current"] = own["current"].map({True: "Current", False: "Former"})
    own.columns = ["CCN", "Facility", "Current?", "Role", "Owner name", "Owner type", "Ownership %", "CMS 'since' date",
                   "First release listing them", "Last release listing them"]
    sheet("Ownership", own, widths={"Facility": 32, "Role": 30, "Owner name": 36})

    # Monthly
    pm = D["pm"].copy()
    pm.insert(1, "Facility", pm["ccn"].map(name_of))
    keep = ["ccn", "Facility", "cms_month", "provider_name", "legal_business_name", "chain_name", "ownership_type"] + \
           [m[1] for m in METRICS] + ["occupancy_pct", "special_focus_status", "abuse_icon",
                                      "provider_changed_ownership_in_last_12_months", "long_stay_qm_rating", "short_stay_qm_rating",
                                      "number_of_citations_from_infection_control_inspections", "rating_cycle_1_standard_survey_health_date",
                                      "source_note"]
    keep = [k for k in keep if k in pm.columns]
    pm = pm[keep]
    lab = {m[1]: m[2] for m in METRICS}
    lab.update({"ccn": "CCN", "cms_month": "CMS release month", "provider_name": "Name at the time",
                "legal_business_name": "Legal business name at the time", "chain_name": "CMS chain",
                "ownership_type": "Ownership type", "occupancy_pct": "Occupancy %", "special_focus_status": "Special Focus status",
                "abuse_icon": "Abuse icon", "provider_changed_ownership_in_last_12_months": "Changed ownership in last 12 months",
                "long_stay_qm_rating": "Long-stay QM rating", "short_stay_qm_rating": "Short-stay QM rating",
                "number_of_citations_from_infection_control_inspections": "Infection control citations",
                "rating_cycle_1_standard_survey_health_date": "Most recent standard inspection", "source_note": "Note"})
    pm.columns = [lab.get(c, c) for c in pm.columns]
    sheet("Monthly", pm.sort_values(["Facility", "CMS release month"]), widths={"Facility": 34, "Name at the time": 30,
                                                                                 "Legal business name at the time": 30})

    # Trend sheets
    for key, title in TREND_SHEETS:
        col = dict((m[0], m[1]) for m in METRICS)[key]
        better = dict((m[0], m[3]) for m in METRICS)[key]
        w = D["pm"].pivot_table(index="ccn", columns="cms_month", values=col, aggfunc="last")
        w = w[sorted(w.columns, reverse=True)]
        w.columns = [month_label(c) for c in w.columns]
        w.insert(0, "Facility", w.index.map(name_of))
        st_ = D["fac"].set_index("ccn")["status"]
        w.insert(1, "Status", w.index.map(lambda c: "Open" if st_.get(c) == "Open" else "Closed"))
        w = w.reset_index().rename(columns={"ccn": "CCN"})
        w = w.sort_values(["Status", "Facility"], ascending=[False, True])
        fmt = {c: "0.00" for c in w.columns[3:]} if "hprd" in key else None
        ws = sheet(title, w, widths={"Facility": 38, "Status": 8}, freeze="D2", fmt=fmt, wrap_header=True)
        if key in ("overall", "health", "staffing", "qm"):
            good_bad(ws, w, list(w.columns[3:]), "stars")
        else:
            n = len(w) + 1
            lo, hi = ("FFF2A8", "BDD7EE") if better == "down" else ("BDD7EE", "FFF2A8")
            rng = f"D2:{get_column_letter(len(w.columns))}{n}"
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="percentile", start_value=10, start_color=lo,
                                                              mid_type="percentile", mid_value=50, mid_color="FFFFFF",
                                                              end_type="percentile", end_value=90, end_color=hi))

    # Citations
    c = D["cit"].copy()
    c.insert(1, "Facility", c["ccn"].map(name_of))
    c = c[["ccn", "Facility", "survey_date", "inspection", "tag", "deficiency_description", "deficiency_category",
           "scope_severity_code", "severity", "scope", "complaint_deficiency", "standard_deficiency", "infection_control_inspection_deficiency",
           "deficiency_corrected", "correction_date", "citation_under_idr", "cms_status", "report_url", "first_seen", "last_seen"]]
    c.columns = ["CCN", "Facility", "Inspection date", "Inspection type", "Tag", "Description", "Category", "Scope/severity",
                 "Severity", "How widespread", "From complaint", "From standard inspection", "Infection control", "Status",
                 "Correction date", "Under dispute (IDR)", "In CMS data now?", "Inspection report (Care Compare)",
                 "First release", "Last release"]
    ws = sheet("Citations", c, widths={"Facility": 32, "Description": 60, "Category": 30, "Severity": 22, "Status": 30,
                                       "How widespread": 22, "Inspection report (Care Compare)": 14})
    ucol = list(c.columns).index("Inspection report (Care Compare)") + 1
    for r in range(2, len(c) + 2):
        cell = ws.cell(r, ucol)
        if cell.value:
            cell.hyperlink = cell.value
            cell.value = "Open report"
            cell.font = Font(name="Helvetica", size=10, underline="single")

    cy = D["cy"].copy()
    cy.insert(1, "Facility", cy["ccn"].map(name_of))
    cy.columns = ["CCN", "Facility", "Year", "Health citations", "Complaint citations", "Harm citations (G-L)",
                  "Immediate jeopardy (J-L)", "Infection control citations", "Fire safety citations"]
    sheet("Citations by Year", cy.sort_values(["Facility", "Year"]), widths={"Facility": 36})

    pen = D["pen"].copy()
    pen.insert(1, "Facility", pen["ccn"].map(name_of))
    pen = pen[["ccn", "Facility", "penalty_date", "penalty_type", "fine_amount", "payment_denial_start_date",
               "payment_denial_length_in_days", "cms_status", "first_seen", "last_seen"]]
    pen.columns = ["CCN", "Facility", "Date", "Type", "Fine amount", "Payment denial start", "Payment denial days",
                   "CMS status", "First release", "Last release"]
    sheet("Fines", pen, widths={"Facility": 36}, fmt={"Fine amount": "$#,##0"})
    py = D["py"].copy()
    py.insert(1, "Facility", py["ccn"].map(name_of))
    py.columns = ["CCN", "Facility", "Year", "Fines", "Fine dollars", "Payment denials"]
    sheet("Fines by Year", py.sort_values(["Facility", "Year"]), widths={"Facility": 36}, fmt={"Fine dollars": "$#,##0"})

    qm = D["qm"].copy()
    meas = D["meas"]
    qm.insert(1, "Facility", qm["ccn"].map(name_of))
    qm.insert(3, "Measure", qm["measure_code"].map(meas["measure_description"]))
    qm = qm[["ccn", "Facility", "measure_code", "Measure", "quarter", "score"]].sort_values(["Facility", "measure_code", "quarter"])
    qm.columns = ["CCN", "Facility", "Measure code", "Measure", "Quarter", "Score"]
    sheet("Quality Measures", qm, widths={"Facility": 32, "Measure": 60})
    qc = D["qc"].copy()
    qc.insert(1, "Facility", qc["ccn"].map(name_of))
    qc.insert(3, "Measure", qc["measure_code"].map(meas["measure_description"]))
    qc = qc[["ccn", "Facility", "measure_code", "Measure", "measure_period", "adjusted_score", "observed_score", "expected_score"]]
    qc.columns = ["CCN", "Facility", "Measure code", "Measure", "Period", "Risk-adjusted score", "Observed", "Expected"]
    sheet("Hospital Measures", qc.sort_values(["Facility", "Measure code", "Period"]), widths={"Facility": 32, "Measure": 60})

    a = D["avg"].copy()
    a = a.sort_values(["cms_month", "area"])
    a.columns = ["CMS release month", "Area"] + [c.replace("_", " ").capitalize() for c in a.columns[2:]]
    sheet("RI vs US", a, freeze="C2")
    r = D["releases"].copy()
    r.columns = ["CMS release month", "Source file", "Loaded", "RI facilities"]
    sheet("Releases", r, freeze="A2", widths={"Source file": 60})
    wb.save(path)


# ------------------------------------------------------------------ dashboard data
def write_dashboard_data(D, path):
    months = D["all_months"]
    idx = {m: i for i, m in enumerate(months)}
    pm = D["pm"]
    keys = [m for m in METRICS]
    series = {}
    for ccn, g in pm.groupby("ccn"):
        s = {}
        for key, col, _, _ in keys:
            arr = [None] * len(months)
            for m, v in zip(g["cms_month"], g[col]):
                if pd.notna(v):
                    arr[idx[m]] = round(float(v), 3)
            if any(x is not None for x in arr):
                s[key] = arr
        sff = [None] * len(months)
        for m, v in zip(g["cms_month"], g.get("special_focus_status", pd.Series(dtype=str))):
            if isinstance(v, str):
                sff[idx[m]] = v
        s["sff"] = sff
        ab = [None] * len(months)
        for m, v in zip(g["cms_month"], g.get("abuse_icon", pd.Series(dtype=str))):
            if v == "Y":
                ab[idx[m]] = 1
        s["abuse"] = ab
        series[ccn] = s

    avg = {}
    amap = {"overall": "overall_rating", "health": "health_inspection_rating", "qm": "qm_rating", "staffing": "staffing_rating",
            "hprd": "reported_total_nurse_staffing_hours_per_resident_per_day", "rn_hprd": "reported_rn_staffing_hours_per_resident_per_day",
            "lpn_hprd": "reported_lpn_staffing_hours_per_resident_per_day", "aide_hprd": "reported_nurse_aide_staffing_hours_per_resident_per_day",
            "wkend_hprd": "total_number_of_nurse_staff_hours_per_resident_per_day_on_the_weekend",
            "cm_hprd": "case_mix_total_nurse_staffing_hours_per_resident_per_day",
            "turnover": "total_nursing_staff_turnover", "rn_turnover": "registered_nurse_turnover",
            "admin_left": "number_of_administrators_who_have_left_the_nursing_home", "cmi": "nursing_case_mix_index",
            "residents": "average_number_of_residents_per_day"}
    for area, g in D["avg"].groupby("area"):
        a = {}
        for k, col in amap.items():
            arr = [None] * len(months)
            for m, v in zip(g["cms_month"], g[col]):
                if m in idx and pd.notna(v):
                    arr[idx[m]] = round(float(v), 3)
            a[k] = arr
        avg[area] = a
    # RI averages computed from facilities where CMS doesn't publish them (e.g. star ratings before 2023)
    for k in ["overall", "health", "qm", "staffing", "hprd", "rn_hprd", "turnover"]:
        col = dict((m[0], m[1]) for m in METRICS)[k]
        mean = pm.groupby("cms_month")[col].mean()
        avg.setdefault("RI_calc", {})[k] = [None if pd.isna(mean.get(m)) else round(float(mean.get(m)), 3) for m in months]
    # statewide counts by year
    cy = D["cy"].groupby("year")[["health_citations", "complaint_citations", "harm_citations", "immediate_jeopardy"]].sum()
    py = D["py"].groupby("year")[["fines", "fine_dollars"]].sum()
    statewide = {"cit_by_year": {y: [int(v) for v in r] for y, r in cy.iterrows()},
                 "fines_by_year": {y: [int(r["fines"]), float(r["fine_dollars"])] for y, r in py.iterrows()},
                 "homes_by_month": [int(pm.loc[pm["cms_month"] == m, "ccn"].nunique()) if m in D["months"] else None for m in months],
                 "beds_by_month": [float(pm.loc[pm["cms_month"] == m, "number_of_certified_beds"].sum()) if m in D["months"] else None for m in months],
                 "residents_by_month": [float(pm.loc[pm["cms_month"] == m, "average_number_of_residents_per_day"].sum()) if m in D["months"] else None for m in months]}

    fac = []
    for _, f in D["fac"].iterrows():
        x = {k[2:]: v for k, v in f.items() if k.startswith("x_") and isinstance(v, str) and v != ""}
        rec = {k: (None if (isinstance(v, float) and np.isnan(v)) else v) for k, v in f.items()
               if not k.startswith("x_") and k not in ("cms_name",)}
        rec["x"] = x
        fac.append(rec)

    tags = {}
    cits = {}
    for ccn, g in D["cit"].groupby("ccn"):
        rows = []
        for _, r in g.iterrows():
            t = r["tag"]
            tags.setdefault(t, [r["deficiency_description"], r["deficiency_category"]])
            rows.append([r["survey_date"], t, r["scope_severity_code"], 1 if r["complaint"] else 0,
                         1 if r["infection_control"] else 0, "F" if r["inspection"] == "Fire safety" else "H",
                         r["deficiency_corrected"] if isinstance(r["deficiency_corrected"], str) else "",
                         1 if r["listed_now"] else 0])
        cits[ccn] = rows
    pens = {}
    for ccn, g in D["pen"].groupby("ccn"):
        pens[ccn] = [[r["penalty_date"], r["penalty_type"], None if pd.isna(r["fine_amount"]) else float(r["fine_amount"]),
                      None if pd.isna(r["payment_denial_length_in_days"]) else int(r["payment_denial_length_in_days"]),
                      r["payment_denial_start_date"] if isinstance(r["payment_denial_start_date"], str) else "",
                      0 if r["cms_status"].startswith("Removed") else (1 if r["cms_status"].startswith("Listed") else 2)]
                     for _, r in g.iterrows()]
    owners = {}
    for ccn, g in D["own"].groupby("ccn"):
        owners[ccn] = [[r["role"], r["owner_name"], r["owner_type"], r["ownership_percentage"] if isinstance(r["ownership_percentage"], str) else "",
                        r["since"] if isinstance(r["since"], str) else "", r["first_seen"], r["last_seen"], 1 if r["current"] else 0]
                       for _, r in g.iterrows()]
    surveys = {}
    tmap = {"Health Standard": "S", "Health Complaint": "C", "Infection Control": "I", "Fire Safety Standard": "F",
            "Fire Safety Complaint": "FC"}
    for ccn, g in D["surveys"].groupby("ccn"):
        surveys[ccn] = [[r["survey_date"], tmap.get(r["type_of_survey"], ""), 1 if r["last_seen"] == D["latest"] else 0]
                        for _, r in g.iterrows()]
    names = {}
    for ccn, g in D["names"].groupby("ccn"):
        names[ccn] = [[r["kind"][0], r["name"], r["from"], r["to"]] for _, r in g.iterrows()]

    qms = {}
    qm = D["qm"][D["qm"]["measure_code"].isin(KEY_QMS)]
    for ccn, g in qm.groupby("ccn"):
        qms[ccn] = {code: [[q, round(float(s), 2)] for q, s in zip(gg["quarter"], gg["score"])]
                    for code, gg in g.sort_values("quarter").groupby("measure_code")}
    qc = D["qc"][D["qc"]["measure_code"].isin(KEY_QMS)].copy()
    qc["end"] = qc["measure_period"].str[-8:-4] + "Q" + ((qc["measure_period"].str[-4:-2].astype(float) - 1) // 3 + 1).astype("Int64").astype(str)
    for ccn, g in qc.groupby("ccn"):
        d_ = qms.setdefault(ccn, {})
        for code, gg in g.sort_values("end").groupby("measure_code"):
            d_[code] = [[e, round(float(s), 2)] for e, s in zip(gg["end"], gg["adjusted_score"])]
    qm_ri = {}
    for code in KEY_QMS:
        rows = {}
        for ccn, d_ in qms.items():
            for q, s in d_.get(code, []):
                rows.setdefault(q, []).append(s)
        qm_ri[code] = [[q, round(sum(v) / len(v), 2)] for q, v in sorted(rows.items())]

    data = {
        "meta": {"latest": D["latest"], "latest_label": month_label(D["latest"]), "built": datetime.now().strftime("%Y-%m-%d"),
                 "first": months[0], "cms_first": D["months"][0]},
        "months": months,
        "metrics": [{"key": k, "label": lbl, "better": b} for k, _, lbl, b in METRICS],
        "facilities": fac, "series": series, "avg": avg, "statewide": statewide,
        "citations": cits, "tags": tags, "penalties": pens, "owners": owners, "names": names,
        "qm_labels": KEY_QMS, "qms": qms, "qm_ri": qm_ri, "surveys": surveys,
        "links": {"pdf": CC_PDF, "cc": CC_PAGE, "pp": PROPUBLICA},
        "harm_summaries": D.get("harm_summaries", {}),
    }
    txt = json.dumps(data, separators=(",", ":"), default=lambda o: None if (isinstance(o, float) and np.isnan(o)) else str(o))
    txt = txt.replace("NaN", "null")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        f.write("window.RINH_DATA=" + txt + ";")
    return len(txt)


def sync_harm_summaries(root, D):
    """Keep harm_summaries.xlsx listing every harm citation (G-L) still on Care Compare, plus any already
    summarized. People type a short plain-language summary in the 'what_happened' column; it shows on the site."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
    path = os.path.join(root, "harm_summaries.xlsx")
    old = {}
    if os.path.exists(path):
        x = pd.read_excel(path, dtype=str).fillna("")
        for _, r in x.iterrows():
            k = (r["ccn"].replace(".0", "").zfill(6), r["inspection_date"][:10], r["tag"])
            old[k] = {"what_happened": r.get("what_happened", ""), "summarized_by": r.get("summarized_by", "")}
    c = D["cit"]
    h = c[(c["inspection"] == "Health") & c["harm"] & (c["listed_now"] | c.apply(
        lambda r: bool(old.get((r["ccn"], r["survey_date"], r["tag"]), {}).get("what_happened")), axis=1))]
    h = h.sort_values(["survey_date", "ccn"], ascending=[False, True])
    rows = []
    for _, r in h.iterrows():
        k = (r["ccn"], r["survey_date"], r["tag"])
        o = old.get(k, {})
        rows.append({"ccn": r["ccn"], "facility": D["name_of"].get(r["ccn"], ""), "inspection_date": r["survey_date"],
                     "tag": r["tag"], "severity": r["scope_severity_code"], "cms_description": r["deficiency_description"],
                     "report_link": r["report_url"], "what_happened": o.get("what_happened", ""),
                     "summarized_by": o.get("summarized_by", "")})
    df = pd.DataFrame(rows)
    wb = Workbook(); ws = wb.active; ws.title = "Harm citations"
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    F = Font(name="Helvetica", size=10); FB = Font(name="Helvetica", size=10, bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = F; cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (6, 8))
    for cell in ws[1]:
        cell.font = FB
    lc = list(df.columns).index("report_link") + 1
    for r_ in range(2, ws.max_row + 1):
        cell = ws.cell(r_, lc)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(name="Helvetica", size=10, underline="single")
    for col, w in zip("ABCDEFGHI", [9, 30, 12, 8, 8, 45, 22, 60, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    n2 = wb.create_sheet("How to fill this in")
    for line in ["Type a short plain-language summary of what happened in the what_happened column, one or two sentences,",
                 "from the official report (report_link). Example: 'Resident fell from a mechanical lift used by one aide instead of two;",
                 "broke a hip and was hospitalized.' Put your initials in summarized_by.",
                 "Leave the other columns alone. Each monthly update adds new harm citations to this list and keeps what you wrote.",
                 "Blank summaries are fine; the site shows a summary only where one is filled in."]:
        n2.append([line])
    n2.column_dimensions["A"].width = 120
    wb.save(path)
    return {f"{r['ccn']}|{r['inspection_date']}|{r['tag']}": r["what_happened"] for r in rows if r["what_happened"].strip()}


def build_all(root):
    d = load(root)
    D = derive(d)
    D["harm_summaries"] = sync_harm_summaries(root, D)
    out = os.path.join(root, "out")
    os.makedirs(out, exist_ok=True)
    write_excel(D, os.path.join(out, "RI_Nursing_Homes_Database.xlsx"))
    n = write_dashboard_data(D, os.path.join(out, "dashboard", "data.js"))
    print(f"Built workbook and dashboard data ({n / 1e6:.1f} MB) through the {month_label(D['latest'])} CMS release.")
    return D


if __name__ == "__main__":
    build_all(os.path.dirname(HERE))
